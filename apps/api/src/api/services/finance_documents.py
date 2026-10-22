"""從核定預算與實際支出產生可下載的財務文件。"""

from collections import OrderedDict
from io import BytesIO
from typing import Protocol
from zipfile import ZIP_DEFLATED, ZipFile

from defusedxml import ElementTree as ET

from api.services.finance import effective_budget_submission_ids

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class _XmlElement(Protocol):
    text: str | None

    def makeelement(self, tag: str, attrib: dict[str, str]) -> "_XmlElement": ...

    def append(self, subelement: "_XmlElement") -> None: ...


def _word(parent: _XmlElement, tag: str, text: str | None = None, **attrs: str) -> _XmlElement:
    element = parent.makeelement(f"{{{W}}}{tag}", {f"{{{W}}}{k}": v for k, v in attrs.items()})
    element.text = text
    parent.append(element)
    return element


def advance_receipt(budget_name: str, expenses: list[dict]) -> bytes:
    """Leave signatures blank: generating a receipt does not confirm a repayment."""
    groups: dict[str, list[tuple[str, int]]] = OrderedDict()
    for expense in expenses:
        recipient = expense["advanced_by_name"]
        rows = groups.setdefault(recipient, [])
        if expense["items"]:
            rows.extend((item["name"], item["amount"]) for item in expense["items"])
        else:
            rows.append((expense["purpose"], expense["total_amount"]))

    # defusedxml exposes parsing/serialization; create child nodes through the safe root.
    document = ET.fromstring(f'<w:document xmlns:w="{W}"/>', forbid_dtd=True)
    body = _word(document, "body")
    for text in (
        "班聯內部實務利用之預先代墊簽收單",
        budget_name,
        "本表供代墊款項核銷及返還使用。確認支出憑證與代墊金額後，由本會返還實際墊付款項；"
        "代墊人收受款項後簽名確認。",
    ):
        _word(_word(_word(body, "p"), "r"), "t", text)
    table = _word(body, "tbl")
    properties = _word(table, "tblPr")
    _word(properties, "tblW", w="9900", type="dxa")
    borders = _word(properties, "tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _word(borders, edge, val="single", sz="6", color="000000")
    grid = _word(table, "tblGrid")
    widths = (1700, 3400, 1400, 1600, 1800)
    for width in widths:
        _word(grid, "gridCol", w=str(width))

    def append_row(values: list[str], merge: str | None = None, header: bool = False) -> None:
        row = _word(table, "tr")
        row_properties = _word(row, "trPr")
        _word(row_properties, "cantSplit")
        _word(row_properties, "trHeight", val="520", hRule="atLeast")
        if header:
            _word(row_properties, "tblHeader")
        for index, (value, width) in enumerate(zip(values, widths, strict=True)):
            cell = _word(row, "tc")
            props = _word(cell, "tcPr")
            _word(props, "tcW", w=str(width), type="dxa")
            _word(props, "vAlign", val="center")
            if merge and index in (0, 3, 4):
                _word(props, "vMerge", val=merge)
            _word(_word(_word(cell, "p"), "r"), "t", value)

    append_row(["代墊人", "代墊項目", "代墊金額", "總金額", "簽名"], header=True)
    for recipient, rows in groups.items():
        total = sum(amount for _, amount in rows)
        for index, (name, amount) in enumerate(rows):
            append_row(
                [
                    recipient if index == 0 else "",
                    name,
                    f"{amount:,}",
                    f"{total:,}" if index == 0 else "",
                    "",
                ],
                merge="restart" if index == 0 else "continue",
            )
    _word(_word(_word(body, "p"), "r"), "t", "簽收日期：民國　　　年　　　月　　　日")
    section = _word(body, "sectPr")
    _word(section, "pgSz", w="11906", h="16838")
    _word(section, "pgMar", top="1000", right="900", bottom="1000", left="900")
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/word/styles.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
        "</Types>"
    )
    root_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/></Relationships>'
    )
    document_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
        'Target="styles.xml"/></Relationships>'
    )
    styles = (
        f'<w:styles xmlns:w="{W}"><w:docDefaults><w:rPrDefault><w:rPr>'
        '<w:rFonts w:ascii="Times New Roman" w:eastAsia="標楷體"/>'
        '<w:sz w:val="24"/></w:rPr></w:rPrDefault></w:docDefaults></w:styles>'
    )
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr(
            "word/document.xml", ET.tostring(document, encoding="utf-8", xml_declaration=True)
        )
        archive.writestr("word/_rels/document.xml.rels", document_rels)
        archive.writestr("word/styles.xml", styles)
    return output.getvalue()


def budget_workbook(detail: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    effective_ids = effective_budget_submission_ids(detail["submissions"])
    nodes = {node["id"]: node for node in detail["nodes"]}
    rows = []
    for allocation in detail["allocations"]:
        if allocation["submission_id"] not in effective_ids:
            continue
        node = nodes.get(allocation["node_id"], {})
        parent = nodes.get(node.get("parent_id"), {})
        rows.append(
            {
                "category": parent.get("name", node.get("name", "未分類")),
                "name": node.get("name", "未命名細項"),
                "quantity": f"{allocation['quantity']:g}{allocation['unit'] or ''}"
                if allocation["quantity"]
                else "1式",
                "unit_price": allocation["unit_price"] or "*",
                "amount": allocation["amount"],
                "note": allocation["note"],
                "is_new": allocation["is_new"],
                "income": False,
            }
        )
    for income in detail["income_items"]:
        if income.submission_id in effective_ids:
            rows.append(
                {
                    "category": income.category,
                    "name": income.name,
                    "quantity": f"{income.quantity:g}{income.unit or ''}"
                    if income.quantity
                    else "1式",
                    "unit_price": income.unit_price or "*",
                    "amount": income.amount,
                    "note": income.note,
                    "is_new": income.is_new,
                    "income": True,
                }
            )
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, selected in (
        ("預算總表", rows),
        ("本次新增", [r for r in rows if r["is_new"]]),
    ):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["項目", "細項", "數量", "單價", "總額(含稅)", "項目總額", "備註"])
        groups: dict[str, list[dict]] = OrderedDict()
        for row in selected:
            groups.setdefault(row["category"], []).append(row)
        for category, entries in groups.items():
            start = sheet.max_row + 1
            balance = sum(r["amount"] * (1 if r["income"] else -1) for r in entries)
            for index, row in enumerate(entries):
                sheet.append(
                    [
                        category if index == 0 else None,
                        row["name"],
                        row["quantity"],
                        row["unit_price"],
                        row["amount"],
                        abs(balance) if index == 0 else None,
                        row["note"],
                    ]
                )
                if row["is_new"]:
                    for cell in sheet[sheet.max_row]:
                        cell.fill = PatternFill("solid", fgColor="FCE5CD")
                sheet.cell(sheet.max_row, 5).fill = PatternFill(
                    "solid", fgColor="D9EAD3" if row["income"] else "F4CCCC"
                )
            end = sheet.max_row
            if end > start:
                sheet.merge_cells(start_row=start, end_row=end, start_column=1, end_column=1)
                sheet.merge_cells(start_row=start, end_row=end, start_column=6, end_column=6)
            sheet.cell(start, 6).fill = PatternFill(
                "solid", fgColor="D9EAD3" if balance >= 0 else "F4CCCC"
            )
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="微軟正黑體", size=11, bold=cell.row == 1)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for index, width in enumerate((20, 32, 12, 12, 16, 16, 48), 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "C2"
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
