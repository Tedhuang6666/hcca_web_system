"""銷售分析 PDF／Excel 報告，資料由伺服器按同一組篩選重新查詢。"""

from __future__ import annotations

import asyncio
import base64
import html
import math
from io import BytesIO

import pygal
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pygal.style import Style
from sqlalchemy.ext.asyncio import AsyncSession

from api.core.clock import TAIPEI
from api.models.activity import Activity
from api.models.school_class import SchoolClass
from api.models.shop import Product, ProductCategory
from api.schemas.shop_analytics import ShopSalesAnalyticsOut, ShopSalesReportCreate, ShopSalesRow
from api.services.official_print import _font_faces, render_print_pdf
from api.services.shop._analytics import sales_analytics

METRICS = [
    ("total_amount", "訂購金額"),
    ("paid_amount", "班聯已確認金額"),
    ("unpaid_amount", "班聯未確認金額"),
    ("order_count", "有效訂單數"),
    ("buyer_count", "不重複訂購人數"),
    ("item_count", "商品件數"),
    ("average_order_amount", "平均客單價（排除零元）"),
    ("median_order_amount", "客單價中位數（排除零元）"),
    ("average_buyer_amount", "每人平均訂購金額"),
    ("average_items_per_order", "平均每單件數"),
    ("discount_amount", "優惠折抵金額"),
    ("promotion_usage_rate", "優惠使用率（%）"),
    ("payment_rate", "款項確認率（%）"),
    ("repeat_buyer_count", "多筆訂單人數"),
    ("repeat_buyer_rate", "多筆訂單人數占比（%）"),
]
DIMENSIONS = [
    ("grades", "年級"),
    ("classes", "班級"),
    ("products", "商品"),
    ("variants", "規格"),
    ("sources", "訂購方式"),
    ("weekdays", "星期"),
    ("hours", "時段"),
    ("weekday_hours", "星期與時段"),
    ("statuses", "訂單狀態（含取消退款）"),
    ("buyer_segments", "訂購頻次"),
    ("amount_bands", "訂單金額區間"),
    ("item_bands", "每單件數區間"),
    ("grade_products", "年級與商品"),
    ("class_products", "班級與商品"),
]
NOTES = (
    "有效訂購金額、件數及人數排除測試、取消與退款；訂單狀態表另外列出取消與退款，金額是原訂單快照，不代表退款實際撥付額。",
    "日期、時段與星期皆採台北時間及訂單建立日期；追加、減購與取消會回算原日期，並非每日異動流水帳。",
    "品項／主題僅統計符合的明細，整單折扣按原價比例與最大餘數法分攤；平均客單價及中位數排除零元訂單。",
    "未確認款項包含已交班代、尚未由班聯確認的款項；不是實際欠款認定。不同商品／班級的訂購人數、訂單數可能重疊，不可直接加總。",
    "多筆訂單依同一帳號在所選範圍的有效訂單數計算，追加原訂單不增加筆數；不能直接解讀為回購率。",
    "星期時段熱區為區間總量，未校正各星期在區間出現的次數。未提供成本與瀏覽數，無法推算利潤或轉換率。",
)


async def create_sales_report(
    session: AsyncSession, request: ShopSalesReportCreate
) -> tuple[bytes, str, str]:
    data = await sales_analytics(session, compare=False, **request.model_dump(exclude={"format"}))
    scope = []
    for field, model, prefix in [
        ("activity_id", Activity, "活動"),
        ("category_id", ProductCategory, "商品主題"),
        ("product_id", Product, "商品"),
        ("class_id", SchoolClass, "班級"),
    ]:
        identifier = getattr(request, field)
        if identifier is not None:
            entity = await session.get(model, identifier)
            label = (
                (entity.label or entity.class_code)
                if isinstance(entity, SchoolClass)
                else entity.name
                if entity
                else "無符合資料"
            )
            scope.append(f"{prefix}：{label}")
    if request.grade is not None:
        scope.append(f"年級：{request.grade}")
    if request.is_paid is not None:
        scope.append("班聯已確認" if request.is_paid else "班聯未確認")
    scope_label = "；".join(scope) or "全部活動、全部商品、全部班級、全部收款狀態"
    content = await asyncio.to_thread(
        render_sales_pdf if request.format == "pdf" else render_sales_excel, data, scope_label
    )
    media_type = (
        "application/pdf"
        if request.format == "pdf"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"銷售分析_{data.current.date_from}_{data.current.date_to}.{request.format}"
    return content, media_type, filename


def _value(value: int | float) -> str:
    return f"{value:,.2f}".rstrip("0").rstrip(".") if isinstance(value, float) else f"{value:,}"


def _chart(data: ShopSalesAnalyticsOut, measure: str, title: str) -> str:
    values = [getattr(row, measure) for row in data.current.daily]
    maximum = max(1, *values)
    chart = pygal.Line(
        width=1000,
        height=320,
        title=title,
        show_dots=False,
        js=[],
        legend_at_bottom=True,
        legend_at_bottom_columns=2,
        y_labels=sorted({0, round(maximum / 2), maximum}),
        range=(0, maximum),
        x_label_rotation=20,
        show_minor_x_labels=False,
        x_labels_major_every=max(1, math.ceil(len(data.current.daily) / 8)),
        style=Style(
            background="white",
            plot_background="white",
            foreground="#242433",
            foreground_strong="#242433",
            foreground_subtle="#dddddd",
            colors=("#856518", "#667085"),
            font_family="Noto Sans CJK TC",
        ),
    )
    chart.x_labels = [row.key for row in data.current.daily]
    chart.add("所選區間", [getattr(row, measure) for row in data.current.daily])
    encoded = base64.b64encode(chart.render()).decode("ascii")
    return (
        f'<img class="chart" src="data:image/svg+xml;base64,{encoded}" alt="{html.escape(title)}">'
    )


def _table(headers: list[str], rows: list[list], *, first_text: bool = True) -> str:
    head = "".join(f"<th>{html.escape(label)}</th>" for label in headers)
    body = "".join(
        "<tr>"
        + "".join(
            f'<td class="{"label" if index == 0 and first_text else ""}">{html.escape(_value(cell) if isinstance(cell, (int, float)) else str(cell))}</td>'
            for index, cell in enumerate(row)
        )
        + "</tr>"
        for row in rows
    )
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body or '<tr><td colspan=99>此範圍沒有資料</td></tr>'}</tbody></table>"


def _row_values(row: ShopSalesRow) -> list:
    return [
        row.label,
        row.order_count,
        row.buyer_count,
        row.item_count,
        row.total_amount,
        row.paid_amount,
        row.total_amount - row.paid_amount,
        row.discount_amount,
    ]


ROW_HEADERS = ["項目", "訂單", "人數", "件數", "訂購金額", "已確認", "未確認", "折抵"]


def render_sales_pdf(data: ShopSalesAnalyticsOut, scope: str) -> bytes:
    esc = html.escape
    current = data.current
    sections = [
        f"<h1>銷售分析報告</h1><p>{esc(scope)}</p><p>資料區間 {current.date_from} 至 {current.date_to}</p><p>產生時間：{data.generated_at.astimezone(TAIPEI):%Y-%m-%d %H:%M:%S}（台北時間）</p>",
        "<h2>分析摘要</h2>",
    ]
    sections.extend(
        f"<h3>{esc(item.title)}</h3><p>{esc(item.detail)}</p>" for item in data.insights
    )
    sections.append('<section class="page"><h2>關鍵指標</h2>')
    sections.append(
        _table(["指標", "數值"], [[label, getattr(current.metrics, key)] for key, label in METRICS])
    )
    sections.append(
        '</section><section class="page"><h2>每日趨勢</h2><p>橫軸為實際日期；完整日期與數值見每日明細。</p>'
    )
    sections.extend(
        _chart(data, key, title)
        for key, title in [
            ("order_count", "每日訂單數"),
            ("total_amount", "每日訂購金額"),
            ("cumulative_amount", "累積訂購金額"),
        ]
    )
    sections.append("</section>")
    for key, title in DIMENSIONS:
        sections.append(f'<section class="page"><h2>{esc(title)}分析</h2>')
        sections.append(_table(ROW_HEADERS, [_row_values(row) for row in getattr(current, key)]))
        sections.append("</section>")
    sections.append('<section class="page"><h2>每日完整明細</h2>')
    sections.append(
        _table(
            ["日期", "訂單", "人數", "件數", "訂購金額", "已確認", "累積金額"],
            [
                [
                    row.key,
                    row.order_count,
                    row.buyer_count,
                    row.item_count,
                    row.total_amount,
                    row.paid_amount,
                    row.cumulative_amount,
                ]
                for row in current.daily
            ],
        )
    )
    sections.append(
        "</section><h2>統計口徑與限制</h2>" + "".join(f"<p>{esc(note)}</p>" for note in NOTES)
    )
    css = """@page { size:A4; margin:17mm 13mm 18mm; @bottom-center { content:'銷售分析報告 · 第 ' counter(page) ' 頁'; font-size:9pt; color:#666; } }
body { font-family:"Noto Sans CJK TC",OfficialKai,sans-serif; font-size:10pt; color:#242433; line-height:1.6; } h1 {font-size:25pt;} h2 {font-size:16pt; margin-top:20pt;} h3 {font-size:11pt;margin-bottom:3pt;} p {margin:4pt 0 9pt;} .page {break-before:page;} .chart {width:100%;height:64mm;object-fit:contain;} table {width:100%; border-collapse:collapse;font-size:8.5pt;table-layout:fixed;} thead {display:table-header-group;} th,td {padding:5pt 3pt;border-bottom:0.5pt solid #ddd;text-align:right;overflow-wrap:anywhere;} th {background:#f5f1e5;} th:first-child,td.label {text-align:left;width:26%;} tr {break-inside:avoid;} h2,h3 {break-after:avoid;}"""
    return render_print_pdf(
        f'<!doctype html><html lang="zh-TW"><meta charset="utf-8"><style>{_font_faces()}{css}</style><body>{"".join(sections)}</body></html>'
    )


def _sheet(workbook: Workbook, name: str, headers: list[str], rows: list[list]):
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if cell.row == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="66501D")
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 45 if index == 1 else 20
        for cell in list(sheet.columns)[index - 1][1:]:
            if any(word in header for word in ("金額", "確認", "折抵")):
                cell.number_format = "#,##0.00"
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def render_sales_excel(data: ShopSalesAnalyticsOut, scope: str) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    period = data.current
    _sheet(
        workbook,
        "報告說明",
        ["項目", "內容"],
        [
            ["統計範圍", scope],
            ["資料區間", f"{period.date_from} 至 {period.date_to}"],
            ["產生時間（台北）", data.generated_at.astimezone(TAIPEI).isoformat()],
            *[[item.title, item.detail] for item in data.insights],
            *[["統計口徑", note] for note in NOTES],
        ],
    )
    _sheet(
        workbook,
        "關鍵指標",
        ["指標", "數值"],
        [[label, getattr(period.metrics, key)] for key, label in METRICS],
    )
    daily = _sheet(
        workbook,
        "每日明細",
        ["日期", "訂單數", "訂購人數", "件數", "訂購金額", "班聯已確認", "累積金額"],
        [
            [
                row.key,
                row.order_count,
                row.buyer_count,
                row.item_count,
                row.total_amount,
                row.paid_amount,
                row.cumulative_amount,
            ]
            for row in period.daily
        ],
    )
    for column, title, anchor in [
        (2, "每日訂單數", "I2"),
        (5, "每日訂購金額", "I18"),
        (7, "累積訂購金額", "I34"),
    ]:
        chart = LineChart()
        chart.title = title
        chart.x_axis.title = "日期"
        chart.add_data(
            Reference(daily, min_col=column, min_row=1, max_row=daily.max_row),
            titles_from_data=True,
        )
        chart.set_categories(Reference(daily, min_col=1, min_row=2, max_row=daily.max_row))
        daily.add_chart(chart, anchor)
    for key, label in DIMENSIONS:
        _sheet(workbook, label, ROW_HEADERS, [_row_values(row) for row in getattr(period, key)])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
