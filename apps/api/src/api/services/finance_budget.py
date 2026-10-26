"""預算申請、編列、報支與公開查詢。"""

from __future__ import annotations

import asyncio
import re
import uuid
from datetime import UTC, date, datetime
from decimal import ROUND_HALF_UP, Decimal

from fastapi import HTTPException
from sqlalchemy import case, delete, exists, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from api.models.finance import (
    BudgetSubmissionKind,
    BudgetSubmissionStatus,
    ChartAccount,
    ExpenseClaimItem,
    ExpenseClaimItemEvidence,
    ExpenseClaimStatus,
    ExpensePaymentMethod,
    ExpensePaymentStatus,
    FinanceAccountType,
    FinanceBudget,
    FinanceBudgetAllocation,
    FinanceBudgetAllocationEvidence,
    FinanceBudgetAllocationRevision,
    FinanceBudgetExpense,
    FinanceBudgetExpenseEvidence,
    FinanceBudgetExpenseItem,
    FinanceBudgetIncomeItem,
    FinanceBudgetNode,
    FinanceBudgetSubmission,
    FiscalPeriod,
    JournalEntry,
    JournalLine,
    JournalStatus,
)
from api.models.org import Org
from api.models.user import User
from api.schemas.finance import (
    BudgetAllocationCreate,
    BudgetAllocationEvidenceCreate,
    BudgetAllocationUpdate,
    BudgetCreate,
    BudgetIncomeItemUpdate,
    BudgetNodeCreate,
    BudgetReview,
    BudgetSubmissionCreate,
    BudgetSubmissionUpdate,
    FinanceBudgetExpenseCreate,
    FinanceBudgetExpenseUpdate,
)
from api.services.finance_budget_plan import ApprovedBudgetPlan
from api.services.finance_core import (
    _claim_item_total,
    create_budget,
    get_budget,
    get_ledger,
    validate_evidence_key,
)


def _import_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _import_decimal(value: object) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        parsed = Decimal(str(value).replace(",", "").strip())
    except (ValueError, ArithmeticError):
        return None
    return parsed if parsed.is_finite() and parsed > 0 else None


def _import_amount(value: object) -> int | None:
    parsed = _import_decimal(value)
    if parsed is None:
        return None
    return int(parsed.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _import_quantity(value: object) -> tuple[Decimal, str]:
    text = _import_text(value)
    if not text:
        return Decimal("1"), "項"
    match = re.match(r"^([0-9]+(?:\.[0-9]+)?)\s*(.*)$", text)
    if not match:
        return Decimal("1"), text[:32] or "項"
    quantity = Decimal(match.group(1))
    return quantity, match.group(2).strip()[:32] or "項"


def _parse_budget_workbook(
    file_bytes: bytes, *, full_snapshot: bool = True
) -> tuple[list[dict], list[dict], list[str]]:
    from io import BytesIO

    from defusedxml import ElementTree
    from defusedxml.common import DefusedXmlException
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("無法讀取 xlsx 預算檔，請確認檔案未損壞") from exc

    loaded_theme = workbook.loaded_theme
    try:
        sheet_name = "預算總表" if full_snapshot else "本次新增"
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active
        rows = list(sheet.iter_rows(values_only=False))
    finally:
        workbook.close()

    theme_colors: list[str | None] = []
    if loaded_theme:
        try:
            theme = ElementTree.fromstring(loaded_theme)
            color_scheme = theme.find(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme"
            )
            if color_scheme is not None:
                for color in color_scheme:
                    definition = next(iter(color), None)
                    theme_colors.append(
                        (definition.attrib.get("val") or definition.attrib.get("lastClr"))
                        if definition is not None
                        else None
                    )
        except (ElementTree.ParseError, DefusedXmlException):
            theme_colors = []

    def is_green_fill(cell: object) -> bool:
        fill = getattr(cell, "fill", None)
        if fill is None or fill.fill_type != "solid":
            return False
        color = fill.fgColor
        rgb = None
        if color.type == "rgb" and color.rgb:
            rgb = color.rgb[-6:]
        elif color.type == "indexed" and color.indexed is not None:
            from openpyxl.styles.colors import COLOR_INDEX

            if 0 <= color.indexed < len(COLOR_INDEX):
                rgb = COLOR_INDEX[color.indexed][-6:]
        elif (
            color.type == "theme"
            and color.theme is not None
            and 0 <= color.theme < len(theme_colors)
        ):
            rgb = theme_colors[color.theme]
        if not rgb or len(rgb) != 6:
            return False
        try:
            red, green, blue = (int(rgb[index : index + 2], 16) for index in (0, 2, 4))
        except ValueError:
            return False
        tint = color.tint or 0
        if tint:
            channels = [red, green, blue]
            for index, channel in enumerate(channels):
                normalized = channel / 255
                channels[index] = round(
                    255 * (normalized * (1 + tint) if tint < 0 else normalized * (1 - tint) + tint)
                )
            red, green, blue = channels
        return green >= 70 and green - red >= 12 and green - blue >= 8

    header_index = None
    columns: dict[str, int] = {}
    for index, row in enumerate(rows[:30]):
        normalized = {
            _import_text(cell.value).replace(" ", ""): column for column, cell in enumerate(row)
        }
        if "項目" in normalized and "細項" in normalized:
            header_index = index
            columns = normalized
            break
    if header_index is None:
        raise ValueError("找不到預算表欄位；至少需要「項目」與「細項」欄")

    category_column = columns["項目"]
    detail_column = columns["細項"]
    quantity_column = columns.get("數量", 3)
    unit_price_column = columns.get("單價", 4)
    amount_column = columns.get("總額(含稅)", columns.get("總額", 5))
    note_column = columns.get("備註", 7)
    imported: list[dict] = []
    imported_income: list[dict] = []
    skipped: list[str] = []
    current_category = ""
    for row_number, row in enumerate(rows[header_index + 1 :], header_index + 2):
        category_cell = row[category_column] if category_column < len(row) else None
        detail_cell = row[detail_column] if detail_column < len(row) else None
        category = _import_text(category_cell.value if category_cell else None)
        detail = _import_text(detail_cell.value if detail_cell else None)
        if category:
            current_category = category
        if not detail:
            continue
        category = category or current_category or "未分類"
        quantity, unit = _import_quantity(
            row[quantity_column].value if quantity_column < len(row) else None
        )
        unit_price = _import_amount(
            row[unit_price_column].value if unit_price_column < len(row) else None
        )
        amount_cell = row[amount_column] if amount_column < len(row) else None
        amount = _import_amount(amount_cell.value if amount_cell else None)
        if amount is None and unit_price is not None:
            amount = int((quantity * unit_price).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        if amount is None:
            skipped.append(f"第 {row_number} 列「{detail}」沒有可用總額")
            continue
        note = _import_text(row[note_column].value if note_column < len(row) else None)
        # Orange marks newly added details; the amount's green/red fill is a separate dimension.
        is_new = any(
            cell is not None
            and getattr(cell, "fill", None) is not None
            and cell.fill.fill_type == "solid"
            and cell.fill.fgColor.type == "rgb"
            and cell.fill.fgColor.rgb[-6:].upper() in {"FCE5CD", "F9CB9C", "F4B183"}
            for cell in (category_cell, detail_cell)
        )
        is_income = is_green_fill(amount_cell)
        if is_income:
            imported_income.append(
                {
                    "category": category[:160],
                    "name": detail[:160],
                    "amount": amount,
                    "is_new": is_new,
                    "note": note or None,
                    "source_row_number": row_number,
                    "quantity": quantity,
                    "unit": unit,
                    "unit_price": unit_price,
                }
            )
            continue
        if unit_price is not None:
            calculated = int((quantity * unit_price).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
            if calculated != amount:
                note = "；".join(filter(None, [note, "單價與總額不一致，保留試算表總額"]))
                unit_price = None
        imported.append(
            {
                "category": category[:160],
                "detail": detail[:160],
                "quantity": quantity,
                "unit": unit,
                "unit_price": unit_price,
                "amount": amount,
                "is_new": is_new,
                "note": note or None,
                "row_number": row_number,
            }
        )
    if not imported and not imported_income:
        raise ValueError("預算表沒有可匯入的明細")
    return imported, imported_income, skipped


async def import_budget_from_xlsx(
    db: AsyncSession,
    ledger_id: uuid.UUID,
    period_id: uuid.UUID,
    name: str,
    title: str | None,
    file_bytes: bytes,
    user_id: uuid.UUID,
    council_approved_on: date | None,
    proposing_org_id: uuid.UUID | None = None,
    budget_id: uuid.UUID | None = None,
    replace_submission_id: uuid.UUID | None = None,
    full_snapshot: bool = True,
) -> tuple[FinanceBudget, FinanceBudgetSubmission, int, int, int, list[str]]:
    rows, income_rows, skipped = await asyncio.to_thread(
        _parse_budget_workbook, file_bytes, full_snapshot=full_snapshot
    )
    ledger = await get_ledger(db, ledger_id)
    org_id = proposing_org_id or ledger.org_id
    if not await db.get(Org, org_id):
        raise HTTPException(400, "提出部門不存在")

    if replace_submission_id and not budget_id:
        raise HTTPException(400, "重新匯入時必須指定共同預算")
    if budget_id:
        budget = await get_budget(db, budget_id)
        if budget.ledger_id != ledger_id or budget.period_id != period_id:
            raise HTTPException(400, "指定的共同預算不屬於此帳本或會計期間")
        if replace_submission_id:
            submission = await _editable_submission(db, replace_submission_id)
            if submission.budget_id != budget.id:
                raise HTTPException(400, "要覆寫的草案不屬於此共同預算")
            allocation_ids = select(FinanceBudgetAllocation.id).where(
                FinanceBudgetAllocation.submission_id == submission.id
            )
            await db.execute(
                delete(FinanceBudgetAllocationEvidence).where(
                    FinanceBudgetAllocationEvidence.allocation_id.in_(allocation_ids)
                )
            )
            await db.execute(
                delete(FinanceBudgetAllocation).where(
                    FinanceBudgetAllocation.submission_id == submission.id
                )
            )
            await db.execute(
                delete(FinanceBudgetIncomeItem).where(
                    FinanceBudgetIncomeItem.submission_id == submission.id
                )
            )
        else:
            submission = await create_budget_submission(
                db,
                budget,
                BudgetSubmissionCreate(
                    kind="supplemental",
                    title=(title or f"{budget.name}（追加匯入）")[:160],
                    note="由 xlsx 預算表匯入",
                ),
                user_id,
            )
    else:
        budget = await create_budget(db, ledger_id, BudgetCreate(period_id=period_id, name=name))
        submission = await create_budget_submission(
            db,
            budget,
            BudgetSubmissionCreate(
                kind="initial",
                title=(title or f"{name}（匯入）")[:160],
                note="由 xlsx 預算表匯入",
            ),
            user_id,
        )

    categories: dict[str, FinanceBudgetNode] = {}
    for sort_order, row in enumerate(rows):
        category = categories.get(row["category"])
        if category is None:
            category = await _get_or_create_import_node(
                db, budget.id, submission.id, row["category"], None, sort_order
            )
            categories[row["category"]] = category
        detail = await _get_or_create_import_node(
            db, budget.id, submission.id, row["detail"], category.id, sort_order
        )
        allocation = await create_budget_allocation(
            db,
            submission.id,
            BudgetAllocationCreate(
                node_id=detail.id,
                amount=row["amount"],
                quantity=row["quantity"],
                unit=row["unit"],
                unit_price=row["unit_price"],
                proposing_org_id=org_id,
                note=row["note"],
            ),
            user_id,
        )
        allocation.is_new = row["is_new"]
    db.add_all(FinanceBudgetIncomeItem(submission_id=submission.id, **row) for row in income_rows)
    await _refresh_budget_node_activity(db, budget.id)
    budget.is_public = True
    submission.status = BudgetSubmissionStatus.APPROVED
    submission.review_note = "依已核定預算表直接匯入公開，未在系統內重複審核。"
    submission.council_approved_on = council_approved_on
    submission.is_full_snapshot = full_snapshot
    submission.created_at = datetime.now(UTC)
    return budget, submission, len(categories), len(rows), len(income_rows), skipped


async def create_budget_submission(
    db: AsyncSession, budget: FinanceBudget, body: BudgetSubmissionCreate, user_id: uuid.UUID
) -> FinanceBudgetSubmission:
    if body.kind.value == "initial":
        existing = await db.scalar(
            select(FinanceBudgetSubmission.id).where(
                FinanceBudgetSubmission.budget_id == budget.id,
                FinanceBudgetSubmission.kind == body.kind,
            )
        )
        if existing:
            raise HTTPException(409, "共同預算只能建立一份初始預算案")
    submission = FinanceBudgetSubmission(
        budget_id=budget.id, created_by_id=user_id, **body.model_dump()
    )
    db.add(submission)
    await db.flush()
    return submission


async def _editable_submission(
    db: AsyncSession, submission_id: uuid.UUID
) -> FinanceBudgetSubmission:
    submission = await db.get(FinanceBudgetSubmission, submission_id)
    if not submission:
        raise HTTPException(404, "預算案不存在")
    if submission.status not in (BudgetSubmissionStatus.DRAFT, BudgetSubmissionStatus.RETURNED):
        raise HTTPException(400, "只有草案或退回的預算案可以編輯")
    return submission


async def update_budget_submission(
    db: AsyncSession, submission_id: uuid.UUID, body: BudgetSubmissionUpdate
) -> FinanceBudgetSubmission:
    submission = await _editable_submission(db, submission_id)
    values = body.model_dump(exclude_unset=True)
    if not values:
        raise HTTPException(422, "請至少修改預算案名稱或備註")
    if "title" in values:
        title = (values["title"] or "").strip()
        if not title:
            raise HTTPException(422, "預算案名稱不可空白")
        values["title"] = title
    for field, value in values.items():
        setattr(submission, field, value)
    await db.flush()
    return submission


async def delete_budget_submission(db: AsyncSession, submission_id: uuid.UUID) -> None:
    submission = await _editable_submission(db, submission_id)
    budget_id = submission.budget_id
    await db.delete(submission)
    await db.flush()
    await _refresh_budget_node_activity(db, budget_id)


async def create_budget_node(
    db: AsyncSession, submission_id: uuid.UUID, body: BudgetNodeCreate
) -> FinanceBudgetNode:
    submission = await _editable_submission(db, submission_id)
    budget = await get_budget(db, submission.budget_id)
    if body.parent_id:
        parent = await db.get(FinanceBudgetNode, body.parent_id)
        if not parent or parent.budget_id != budget.id:
            raise HTTPException(400, "上層預算條目不屬於此預算")
    node = FinanceBudgetNode(budget_id=budget.id, **body.model_dump())
    db.add(node)
    await db.flush()
    return node


async def _get_or_create_import_node(
    db: AsyncSession,
    budget_id: uuid.UUID,
    submission_id: uuid.UUID,
    name: str,
    parent_id: uuid.UUID | None,
    sort_order: int,
) -> FinanceBudgetNode:
    parent_clause = (
        FinanceBudgetNode.parent_id.is_(None)
        if parent_id is None
        else FinanceBudgetNode.parent_id == parent_id
    )
    existing = await db.scalar(
        select(FinanceBudgetNode).where(
            FinanceBudgetNode.budget_id == budget_id,
            parent_clause,
            FinanceBudgetNode.name == name,
        )
    )
    if existing:
        existing.is_active = True
        existing.sort_order = min(existing.sort_order, sort_order)
        return existing
    return await create_budget_node(
        db,
        submission_id,
        BudgetNodeCreate(parent_id=parent_id, name=name, sort_order=sort_order),
    )


async def _refresh_budget_node_activity(db: AsyncSession, budget_id: uuid.UUID) -> None:
    nodes = list(
        (
            await db.execute(
                select(FinanceBudgetNode).where(FinanceBudgetNode.budget_id == budget_id)
            )
        ).scalars()
    )
    node_by_id = {node.id: node for node in nodes}
    active_ids = set(
        (
            await db.execute(
                select(FinanceBudgetAllocation.node_id)
                .join(FinanceBudgetSubmission)
                .where(FinanceBudgetSubmission.budget_id == budget_id)
            )
        ).scalars()
    )
    for node_id in list(active_ids):
        current = node_by_id.get(node_id)
        while current and current.parent_id:
            active_ids.add(current.parent_id)
            current = node_by_id.get(current.parent_id)
    for node in nodes:
        node.is_active = node.id in active_ids
    await db.flush()


async def create_budget_allocation(
    db: AsyncSession,
    submission_id: uuid.UUID,
    body: BudgetAllocationCreate,
    user_id: uuid.UUID,
) -> FinanceBudgetAllocation:
    submission = await _editable_submission(db, submission_id)
    node = await db.get(FinanceBudgetNode, body.node_id)
    if not node or node.budget_id != submission.budget_id:
        raise HTTPException(400, "預算條目不屬於此預算案")
    child = await db.scalar(
        select(FinanceBudgetNode.id).where(FinanceBudgetNode.parent_id == node.id).limit(1)
    )
    if child:
        raise HTTPException(400, "只有最末層預算條目可以配置金額")
    allocation = FinanceBudgetAllocation(
        submission_id=submission.id, proposed_by_id=user_id, **body.model_dump()
    )
    db.add(allocation)
    await db.flush()
    return allocation


async def update_budget_draft_allocation(
    db: AsyncSession,
    submission_id: uuid.UUID,
    allocation_id: uuid.UUID,
    body: BudgetAllocationCreate,
) -> FinanceBudgetAllocation:
    submission = await _editable_submission(db, submission_id)
    allocation = await db.get(FinanceBudgetAllocation, allocation_id)
    if not allocation or allocation.submission_id != submission.id:
        raise HTTPException(404, "預算配置不存在")
    node = await db.get(FinanceBudgetNode, body.node_id)
    if not node or node.budget_id != submission.budget_id:
        raise HTTPException(400, "預算條目不屬於此預算案")
    child = await db.scalar(
        select(FinanceBudgetNode.id).where(FinanceBudgetNode.parent_id == node.id).limit(1)
    )
    if child:
        raise HTTPException(400, "只有最末層預算條目可以配置金額")
    allocation.node_id = body.node_id
    assert body.amount is not None
    allocation.amount = body.amount
    allocation.quantity = body.quantity
    allocation.unit = body.unit
    allocation.unit_price = body.unit_price
    allocation.proposing_org_id = body.proposing_org_id
    allocation.note = body.note
    await db.flush()
    return allocation


async def delete_budget_draft_allocation(
    db: AsyncSession, submission_id: uuid.UUID, allocation_id: uuid.UUID
) -> None:
    submission = await _editable_submission(db, submission_id)
    allocation = await db.get(FinanceBudgetAllocation, allocation_id)
    if not allocation or allocation.submission_id != submission.id:
        raise HTTPException(404, "預算配置不存在")
    budget_id = submission.budget_id
    await db.delete(allocation)
    await db.flush()
    await _refresh_budget_node_activity(db, budget_id)


async def submit_budget_submission(
    db: AsyncSession, submission_id: uuid.UUID
) -> FinanceBudgetSubmission:
    submission = await _editable_submission(db, submission_id)
    count = await db.scalar(
        select(func.count())
        .select_from(FinanceBudgetAllocation)
        .where(FinanceBudgetAllocation.submission_id == submission.id)
    )
    income_count = await db.scalar(
        select(func.count())
        .select_from(FinanceBudgetIncomeItem)
        .where(FinanceBudgetIncomeItem.submission_id == submission.id)
    )
    if not count and not income_count:
        raise HTTPException(400, "預算案至少需要一筆收入或支出明細")
    submission.status = BudgetSubmissionStatus.SUBMITTED
    submission.submitted_at = datetime.now(UTC)
    await db.flush()
    return submission


async def review_budget_submission(
    db: AsyncSession, submission_id: uuid.UUID, body: BudgetReview, reviewer_id: uuid.UUID
) -> FinanceBudgetSubmission:
    submission = await db.get(FinanceBudgetSubmission, submission_id)
    if not submission:
        raise HTTPException(404, "預算案不存在")
    if submission.status != BudgetSubmissionStatus.SUBMITTED:
        raise HTTPException(400, "只有待審預算案可以審核")
    submission.status = body.status
    submission.review_note = body.note
    submission.reviewed_by_id = reviewer_id
    submission.reviewed_at = datetime.now(UTC)
    submission.council_approved_on = body.council_approved_on
    if body.status in {BudgetSubmissionStatus.APPROVED, BudgetSubmissionStatus.REJECTED}:
        submission.is_council_review_public = False
    await db.flush()
    return submission


async def update_budget_allocation(
    db: AsyncSession,
    allocation_id: uuid.UUID,
    body: BudgetAllocationUpdate,
    user_id: uuid.UUID,
) -> FinanceBudgetAllocation:
    allocation = await db.get(FinanceBudgetAllocation, allocation_id)
    if not allocation:
        raise HTTPException(404, "預算配置不存在")
    submission = await db.get(FinanceBudgetSubmission, allocation.submission_id)
    if not submission or submission.status != BudgetSubmissionStatus.APPROVED:
        raise HTTPException(400, "只有已核准預算可以直接修正")
    budget = await get_budget(db, submission.budget_id)
    node_id = body.node_id or allocation.node_id
    node = await db.get(FinanceBudgetNode, node_id)
    if not node or node.budget_id != budget.id:
        raise HTTPException(400, "預算條目不屬於此預算案")
    child = await db.scalar(
        select(FinanceBudgetNode.id).where(FinanceBudgetNode.parent_id == node.id).limit(1)
    )
    if child:
        raise HTTPException(400, "只有最末層預算條目可以配置金額")

    proposing_org_id = body.proposing_org_id or allocation.proposing_org_id
    if not await db.get(Org, proposing_org_id):
        raise HTTPException(400, "提出部門不存在")

    changed_fields = body.model_dump(exclude_unset=True, exclude={"reason"})
    if "proposing_org_id" in changed_fields and changed_fields["proposing_org_id"] is None:
        raise HTTPException(422, "請選擇提出部門")
    quantity_details_changed = bool({"quantity", "unit", "unit_price"} & changed_fields.keys())
    if quantity_details_changed:
        quantity = changed_fields.get("quantity", allocation.quantity)
        unit = (changed_fields.get("unit", allocation.unit) or "").strip()
        unit_price = changed_fields.get("unit_price", allocation.unit_price)
        if quantity is None or not unit:
            raise HTTPException(422, "填寫數量時，必須同時填寫單位")
        if unit_price is not None:
            calculated = int(
                (Decimal(quantity) * unit_price).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
            )
            if body.amount is not None and body.amount != calculated:
                raise HTTPException(422, "預算總額必須等於數量乘以單價")
            changed_fields["amount"] = calculated
    if not changed_fields:
        raise HTTPException(422, "請至少修改一個預算欄位")

    previous_details = {
        "node_id": str(allocation.node_id),
        "quantity": float(allocation.quantity) if allocation.quantity is not None else None,
        "unit": allocation.unit,
        "unit_price": allocation.unit_price,
        "amount": allocation.amount,
        "note": allocation.note,
        "proposing_org_id": str(allocation.proposing_org_id),
    }
    next_details = {**previous_details}
    for field, value in changed_fields.items():
        setattr(allocation, field, value)
        next_details[field] = (
            str(value)
            if isinstance(value, uuid.UUID)
            else float(value)
            if isinstance(value, Decimal)
            else value
        )
    next_details["node_id"] = str(next_details["node_id"])
    db.add(
        FinanceBudgetAllocationRevision(
            allocation_id=allocation.id,
            previous_amount=previous_details["amount"],
            next_amount=next_details["amount"],
            reason=body.reason,
            changed_by_id=user_id,
            previous_details=previous_details,
            next_details=next_details,
        )
    )
    await db.flush()
    return allocation


async def add_budget_allocation_evidence(
    db: AsyncSession,
    allocation_id: uuid.UUID,
    body: BudgetAllocationEvidenceCreate,
    user_id: uuid.UUID,
) -> FinanceBudgetAllocationEvidence:
    allocation = await db.get(FinanceBudgetAllocation, allocation_id)
    if not allocation:
        raise HTTPException(404, "預算明細不存在")
    submission = await db.get(FinanceBudgetSubmission, allocation.submission_id)
    if not submission:
        raise HTTPException(404, "預算案不存在")
    budget = await get_budget(db, submission.budget_id)
    validate_evidence_key(body.storage_key, budget.ledger_id)
    evidence = FinanceBudgetAllocationEvidence(
        allocation_id=allocation.id,
        uploaded_by_id=user_id,
        **body.model_dump(),
    )
    db.add(evidence)
    await db.flush()
    return evidence


async def _budget_expense_details(db: AsyncSession, budget_id: uuid.UUID) -> list[dict]:
    rows = (
        await db.execute(
            select(FinanceBudgetExpense, FinanceBudgetAllocation, FinanceBudgetNode, Org, User)
            .join(
                FinanceBudgetAllocation,
                FinanceBudgetAllocation.id == FinanceBudgetExpense.allocation_id,
            )
            .join(FinanceBudgetNode, FinanceBudgetNode.id == FinanceBudgetAllocation.node_id)
            .join(
                Org,
                Org.id
                == func.coalesce(
                    FinanceBudgetExpense.department_org_id,
                    FinanceBudgetAllocation.proposing_org_id,
                ),
            )
            .outerjoin(User, User.id == FinanceBudgetExpense.recorded_by_id)
            .where(FinanceBudgetExpense.budget_id == budget_id)
            .order_by(
                FinanceBudgetExpense.entry_date.desc(), FinanceBudgetExpense.created_at.desc()
            )
        )
    ).all()
    expense_ids = [expense.id for expense, _, _, _, _ in rows]
    item_rows = (
        list(
            (
                await db.execute(
                    select(FinanceBudgetExpenseItem)
                    .where(FinanceBudgetExpenseItem.expense_id.in_(expense_ids))
                    .order_by(
                        FinanceBudgetExpenseItem.sort_order,
                        FinanceBudgetExpenseItem.created_at,
                    )
                )
            ).scalars()
        )
        if expense_ids
        else []
    )
    evidence_rows = (
        list(
            (
                await db.execute(
                    select(FinanceBudgetExpenseEvidence)
                    .where(FinanceBudgetExpenseEvidence.expense_id.in_(expense_ids))
                    .order_by(FinanceBudgetExpenseEvidence.created_at)
                )
            ).scalars()
        )
        if expense_ids
        else []
    )
    items_by_expense: dict[uuid.UUID, list[FinanceBudgetExpenseItem]] = {}
    for item in item_rows:
        items_by_expense.setdefault(item.expense_id, []).append(item)
    evidence_by_expense: dict[uuid.UUID, list[FinanceBudgetExpenseEvidence]] = {}
    for evidence in evidence_rows:
        evidence_by_expense.setdefault(evidence.expense_id, []).append(evidence)

    return [
        {
            "id": expense.id,
            "budget_id": expense.budget_id,
            "allocation_id": allocation.id,
            "department_org_id": expense.department_org_id,
            "department_name": department.name,
            "allocation_node_id": node.id,
            "allocation_name": node.name,
            "entry_date": expense.entry_date,
            "operator_name": operator.display_name if operator else "未知操作人",
            "purpose": expense.purpose,
            "advanced_by_name": expense.advanced_by_name,
            "total_amount": expense.total_amount,
            "note": expense.note,
            "items": [
                {
                    "id": item.id,
                    "name": item.name,
                    "unit_price": item.unit_price,
                    "tax_rate": item.tax_rate,
                    "quantity": item.quantity,
                    "unit": item.unit,
                    "amount": _claim_item_total(item.unit_price, item.tax_rate, item.quantity),
                }
                for item in items_by_expense.get(expense.id, [])
            ],
            "evidence": [
                {
                    "id": evidence.id,
                    "storage_key": evidence.storage_key,
                    "filename": evidence.filename,
                    "content_type": evidence.content_type,
                    "file_size": evidence.file_size,
                    "uploaded_at": evidence.created_at,
                    "url": (
                        f"/finance/budgets/{budget_id}/expenses/{expense.id}/evidence/{evidence.id}"
                    ),
                }
                for evidence in evidence_by_expense.get(expense.id, [])
            ],
            "created_at": expense.created_at,
        }
        for expense, allocation, node, department, operator in rows
    ]


async def create_budget_expense(
    db: AsyncSession,
    budget_id: uuid.UUID,
    body: FinanceBudgetExpenseCreate,
    user_id: uuid.UUID,
) -> FinanceBudgetExpense:
    budget = await get_budget(db, budget_id)
    period = await db.get(FiscalPeriod, budget.period_id)
    if not period or not period.starts_on <= body.entry_date <= period.ends_on:
        raise HTTPException(400, "支出日期不在此預算的會計期間內")

    approved_initial = await db.scalar(
        select(FinanceBudgetSubmission.id).where(
            FinanceBudgetSubmission.budget_id == budget.id,
            FinanceBudgetSubmission.kind == BudgetSubmissionKind.INITIAL,
            FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
        )
    )
    if not approved_initial:
        raise HTTPException(400, "請選擇已有核准初始預算的預算案")
    allocation = await db.scalar(
        select(FinanceBudgetAllocation)
        .join(FinanceBudgetSubmission)
        .where(
            FinanceBudgetAllocation.id == body.allocation_id,
            FinanceBudgetSubmission.budget_id == budget.id,
            FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
        )
    )
    if not allocation:
        raise HTTPException(400, "支出必須對應此預算中已核准的明細")
    plan = await ApprovedBudgetPlan.load(db, budget.id)
    if not plan.permits_expense(allocation):
        raise HTTPException(400, "新增支出只能選擇目前核定總表的明細")

    department_org_id = body.department_org_id or allocation.proposing_org_id
    if not await db.get(Org, department_org_id):
        raise HTTPException(400, "提出部門不存在")

    for evidence in body.evidence:
        validate_evidence_key(evidence.storage_key, budget.ledger_id)

    expense = FinanceBudgetExpense(
        budget_id=budget.id,
        allocation_id=allocation.id,
        department_org_id=department_org_id,
        entry_date=body.entry_date,
        purpose=body.purpose,
        advanced_by_name=body.advanced_by_name,
        total_amount=(
            sum(
                _claim_item_total(item.unit_price, item.tax_rate, item.quantity)
                for item in body.items
            )
            if body.items
            else body.total_amount or 0
        ),
        note=body.note,
        recorded_by_id=user_id,
    )
    db.add(expense)
    await db.flush()
    db.add_all(
        [
            FinanceBudgetExpenseItem(
                expense_id=expense.id,
                sort_order=index,
                name=item.name,
                unit_price=item.unit_price,
                tax_rate=item.tax_rate,
                quantity=item.quantity,
                unit=item.unit,
            )
            for index, item in enumerate(body.items)
        ]
    )
    db.add_all(
        [
            FinanceBudgetExpenseEvidence(
                expense_id=expense.id,
                uploaded_by_id=user_id,
                **evidence.model_dump(),
            )
            for evidence in body.evidence
        ]
    )
    await db.flush()
    return expense


async def update_budget_expense(
    db: AsyncSession,
    budget_id: uuid.UUID,
    expense_id: uuid.UUID,
    body: FinanceBudgetExpenseUpdate,
    user_id: uuid.UUID,
) -> FinanceBudgetExpense:
    budget = await get_budget(db, budget_id)
    expense = await db.get(FinanceBudgetExpense, expense_id)
    if not expense or expense.budget_id != budget.id:
        raise HTTPException(404, "預算支出不存在")

    period = await db.get(FiscalPeriod, budget.period_id)
    if not period or not period.starts_on <= body.entry_date <= period.ends_on:
        raise HTTPException(400, "支出日期不在此預算的會計期間內")

    allocation = await db.scalar(
        select(FinanceBudgetAllocation)
        .join(FinanceBudgetSubmission)
        .where(
            FinanceBudgetAllocation.id == body.allocation_id,
            FinanceBudgetSubmission.budget_id == budget.id,
            FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
        )
    )
    if not allocation:
        raise HTTPException(400, "支出必須對應此預算中已核准的明細")
    plan = await ApprovedBudgetPlan.load(db, budget.id)
    if not plan.permits_expense(allocation, original_allocation_id=expense.allocation_id):
        raise HTTPException(400, "只能選擇目前核定明細或此支出原本的歷史明細")

    department_org_id = body.department_org_id or allocation.proposing_org_id
    if not await db.get(Org, department_org_id):
        raise HTTPException(400, "提出部門不存在")
    for evidence in body.evidence:
        validate_evidence_key(evidence.storage_key, budget.ledger_id)

    expense.allocation_id = allocation.id
    expense.department_org_id = department_org_id
    expense.entry_date = body.entry_date
    expense.purpose = body.purpose
    expense.advanced_by_name = body.advanced_by_name
    expense.total_amount = (
        sum(_claim_item_total(item.unit_price, item.tax_rate, item.quantity) for item in body.items)
        if body.items
        else body.total_amount or 0
    )
    expense.note = body.note

    await db.execute(
        delete(FinanceBudgetExpenseItem).where(FinanceBudgetExpenseItem.expense_id == expense.id)
    )
    db.add_all(
        [
            FinanceBudgetExpenseItem(
                expense_id=expense.id,
                sort_order=index,
                name=item.name,
                unit_price=item.unit_price,
                tax_rate=item.tax_rate,
                quantity=item.quantity,
                unit=item.unit,
            )
            for index, item in enumerate(body.items)
        ]
    )
    db.add_all(
        [
            FinanceBudgetExpenseEvidence(
                expense_id=expense.id,
                uploaded_by_id=user_id,
                **evidence.model_dump(),
            )
            for evidence in body.evidence
        ]
    )
    await db.flush()
    return expense


async def budget_advance_receipt_expenses(
    db: AsyncSession, budget_id: uuid.UUID, expense_ids: list[uuid.UUID]
) -> list[dict]:
    await get_budget(db, budget_id)
    expenses = await _budget_expense_details(db, budget_id)
    selected_ids = set(expense_ids)
    selected = [expense for expense in expenses if expense["id"] in selected_ids]
    if len(selected) != len(selected_ids):
        raise HTTPException(422, "所選支出不屬於此預算或已不存在，請重新選取")
    if any(not expense["advanced_by_name"] for expense in selected):
        raise HTTPException(422, "請先為所選支出填寫代墊人")
    return selected


async def update_budget_income_item(
    db: AsyncSession, budget_id: uuid.UUID, income_id: uuid.UUID, body: BudgetIncomeItemUpdate
) -> tuple[FinanceBudgetIncomeItem, int]:
    item = await db.get(FinanceBudgetIncomeItem, income_id)
    submission = await db.get(FinanceBudgetSubmission, item.submission_id) if item else None
    if not item or not submission or submission.budget_id != budget_id:
        raise HTTPException(404, "預算收入不存在")
    plan = await ApprovedBudgetPlan.load(db, budget_id)
    if item.submission_id not in plan.submission_ids:
        raise HTTPException(400, "只能編輯目前核定總表的收入")
    previous = item.amount
    for key, value in body.model_dump(exclude_unset=True).items():
        setattr(item, key, value)
    if item.quantity and item.unit_price and round(item.quantity * item.unit_price) != item.amount:
        item.unit_price = None
    await db.flush()
    return item, previous


async def budget_detail(db: AsyncSession, budget: FinanceBudget) -> dict:
    submissions = list(
        (
            await db.execute(
                select(FinanceBudgetSubmission)
                .where(FinanceBudgetSubmission.budget_id == budget.id)
                .order_by(FinanceBudgetSubmission.created_at)
            )
        ).scalars()
    )
    nodes = list(
        (
            await db.execute(
                select(FinanceBudgetNode)
                .where(FinanceBudgetNode.budget_id == budget.id, FinanceBudgetNode.is_active)
                .order_by(FinanceBudgetNode.sort_order, FinanceBudgetNode.name)
            )
        ).scalars()
    )
    allocations = list(
        (
            await db.execute(
                select(FinanceBudgetAllocation)
                .join(FinanceBudgetSubmission)
                .where(FinanceBudgetSubmission.budget_id == budget.id)
            )
        ).scalars()
    )
    income_items = list(
        (
            await db.execute(
                select(FinanceBudgetIncomeItem)
                .join(FinanceBudgetSubmission)
                .where(FinanceBudgetSubmission.budget_id == budget.id)
                .order_by(
                    FinanceBudgetIncomeItem.submission_id,
                    FinanceBudgetIncomeItem.source_row_number,
                )
            )
        ).scalars()
    )
    allocation_ids = [allocation.id for allocation in allocations]
    evidence_rows = (
        list(
            (
                await db.execute(
                    select(FinanceBudgetAllocationEvidence)
                    .where(FinanceBudgetAllocationEvidence.allocation_id.in_(allocation_ids))
                    .order_by(FinanceBudgetAllocationEvidence.created_at)
                )
            ).scalars()
        )
        if allocation_ids
        else []
    )
    evidence_by_allocation: dict[uuid.UUID, list[FinanceBudgetAllocationEvidence]] = {}
    for evidence in evidence_rows:
        evidence_by_allocation.setdefault(evidence.allocation_id, []).append(evidence)
    plan = ApprovedBudgetPlan(submissions)
    totals = plan.allocated_amounts(allocations)
    used_rows = (
        await db.execute(
            select(
                ExpenseClaimItem.budget_node_id,
                func.coalesce(
                    func.sum(
                        func.round(
                            ExpenseClaimItem.unit_price
                            * ((100 + ExpenseClaimItem.tax_rate) / 100)
                            * ExpenseClaimItem.quantity
                        )
                    ),
                    0,
                ),
            )
            .join(JournalEntry, JournalEntry.id == ExpenseClaimItem.journal_entry_id)
            .where(
                ExpenseClaimItem.budget_node_id.is_not(None),
                JournalEntry.status == JournalStatus.POSTED,
                JournalEntry.ledger_id == budget.ledger_id,
            )
            .group_by(ExpenseClaimItem.budget_node_id)
        )
    ).all()
    used = {row[0]: int(row[1]) for row in used_rows if row[0] is not None}
    direct_used_rows = (
        await db.execute(
            select(FinanceBudgetAllocation.node_id, func.sum(FinanceBudgetExpense.total_amount))
            .join(
                FinanceBudgetExpense,
                FinanceBudgetExpense.allocation_id == FinanceBudgetAllocation.id,
            )
            .join(
                FinanceBudgetSubmission,
                FinanceBudgetSubmission.id == FinanceBudgetAllocation.submission_id,
            )
            .where(
                FinanceBudgetExpense.budget_id == budget.id,
                FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
            )
            .group_by(FinanceBudgetAllocation.node_id)
        )
    ).all()
    for node_id, amount in direct_used_rows:
        used[node_id] = used.get(node_id, 0) + int(amount or 0)
    expenses = await _budget_expense_details(db, budget.id)
    return {
        "id": budget.id,
        "ledger_id": budget.ledger_id,
        "period_id": budget.period_id,
        "name": budget.name,
        "is_public": budget.is_public,
        "submissions": submissions,
        "allocations": [
            {
                "id": allocation.id,
                "submission_id": allocation.submission_id,
                "node_id": allocation.node_id,
                "amount": allocation.amount,
                "is_new": allocation.is_new,
                "quantity": allocation.quantity,
                "unit": allocation.unit,
                "unit_price": allocation.unit_price,
                "note": allocation.note,
                "proposed_by_id": allocation.proposed_by_id,
                "proposing_org_id": allocation.proposing_org_id,
                "evidence": [
                    {
                        "id": evidence.id,
                        "storage_key": evidence.storage_key,
                        "filename": evidence.filename,
                        "content_type": evidence.content_type,
                        "file_size": evidence.file_size,
                        "note": evidence.note,
                        "uploaded_at": evidence.created_at,
                        "url": f"/finance/budget-evidence/{evidence.id}",
                    }
                    for evidence in evidence_by_allocation.get(allocation.id, [])
                ],
            }
            for allocation in allocations
        ],
        "income_items": income_items,
        "expenses": expenses,
        "current_plan": plan.project(allocations, nodes, income_items, expenses),
        "nodes": [
            {
                "id": node.id,
                "budget_id": node.budget_id,
                "parent_id": node.parent_id,
                "name": node.name,
                "sort_order": node.sort_order,
                "allocated_amount": totals.get(node.id, 0),
                "used_amount": used.get(node.id, 0),
                "remaining_amount": totals.get(node.id, 0) - used.get(node.id, 0),
            }
            for node in nodes
        ],
    }


async def set_budget_publication(
    db: AsyncSession, budget: FinanceBudget, is_public: bool
) -> FinanceBudget:
    if is_public:
        approved_initial = await db.scalar(
            select(FinanceBudgetSubmission.id).where(
                FinanceBudgetSubmission.budget_id == budget.id,
                FinanceBudgetSubmission.kind == "initial",
                FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
            )
        )
        if not approved_initial:
            raise HTTPException(400, "初始預算案核准後才能開放議員檢視")
    budget.is_public = is_public
    await db.flush()
    return budget


async def set_submission_council_review_publication(
    db: AsyncSession, submission: FinanceBudgetSubmission, is_public: bool
) -> FinanceBudgetSubmission:
    if is_public and submission.status not in {
        BudgetSubmissionStatus.DRAFT,
        BudgetSubmissionStatus.SUBMITTED,
        BudgetSubmissionStatus.RETURNED,
    }:
        raise HTTPException(400, "只有草案、待審或退回補正的預算案可以開放議員審理")
    submission.is_council_review_public = is_public
    await db.flush()
    return submission


async def list_public_budgets(
    db: AsyncSession,
) -> list[tuple[FinanceBudget, FiscalPeriod, FinanceBudgetSubmission | None]]:
    approved_initial = exists().where(
        FinanceBudgetSubmission.budget_id == FinanceBudget.id,
        FinanceBudgetSubmission.kind == "initial",
        FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
    )
    approved = list(
        (
            await db.execute(
                select(FinanceBudget, FiscalPeriod)
                .join(FiscalPeriod, FiscalPeriod.id == FinanceBudget.period_id)
                .where(FinanceBudget.is_public, approved_initial)
                .order_by(FiscalPeriod.starts_on.desc(), FinanceBudget.name)
            )
        ).all()
    )
    council_review = list(
        (
            await db.execute(
                select(FinanceBudget, FiscalPeriod, FinanceBudgetSubmission)
                .join(FiscalPeriod, FiscalPeriod.id == FinanceBudget.period_id)
                .join(
                    FinanceBudgetSubmission, FinanceBudgetSubmission.budget_id == FinanceBudget.id
                )
                .where(
                    FinanceBudgetSubmission.is_council_review_public,
                    FinanceBudgetSubmission.status.in_(
                        [
                            BudgetSubmissionStatus.DRAFT,
                            BudgetSubmissionStatus.SUBMITTED,
                            BudgetSubmissionStatus.RETURNED,
                        ]
                    ),
                )
                .order_by(
                    FiscalPeriod.starts_on.desc(),
                    FinanceBudget.name,
                    FinanceBudgetSubmission.created_at,
                )
            )
        ).all()
    )
    return [
        *((budget, period, None) for budget, period in approved),
        *((budget, period, submission) for budget, period, submission in council_review),
    ]


async def list_public_expenses(db: AsyncSession, limit: int = 10) -> list[dict]:
    """公開預算的逐項支出紀錄，不回傳承辦人或內部備註。"""
    approved_initial = (
        exists()
        .where(
            FinanceBudgetSubmission.budget_id == FinanceBudget.id,
            FinanceBudgetSubmission.kind == BudgetSubmissionKind.INITIAL,
            FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
        )
        .correlate(FinanceBudget)
    )
    rows = (
        await db.execute(
            select(ExpenseClaimItem, JournalEntry, FinanceBudget, FinanceBudgetNode)
            .join(JournalEntry, JournalEntry.id == ExpenseClaimItem.journal_entry_id)
            .join(FinanceBudgetNode, FinanceBudgetNode.id == ExpenseClaimItem.budget_node_id)
            .join(FinanceBudget, FinanceBudget.id == FinanceBudgetNode.budget_id)
            .where(
                JournalEntry.source_type == "expense_claim",
                JournalEntry.status == JournalStatus.POSTED,
                JournalEntry.claim_status.in_(
                    [ExpenseClaimStatus.APPROVED, ExpenseClaimStatus.COMPLETED]
                ),
                FinanceBudget.is_public.is_(True),
                FinanceBudget.ledger_id == JournalEntry.ledger_id,
                approved_initial,
            )
            .order_by(JournalEntry.entry_date.desc(), JournalEntry.created_at.desc())
            .limit(limit)
        )
    ).all()

    public_expenses: list[tuple[date, datetime, dict]] = []
    for item, entry, budget, node in rows:
        if entry.claim_status == ExpenseClaimStatus.COMPLETED:
            public_status = "spent"
        elif (
            entry.payment_method == ExpensePaymentMethod.ADVANCE
            and entry.payment_status != ExpensePaymentStatus.ADVANCE_REIMBURSED
        ):
            public_status = "awaiting_reimbursement"
        else:
            public_status = "pending"
        public_expenses.append(
            (
                entry.entry_date,
                entry.created_at,
                {
                    "id": item.id,
                    "claim_id": entry.id,
                    "budget_id": budget.id,
                    "entry_date": entry.entry_date,
                    "purpose": re.sub(r"^報帳｜|（\d+ 項）$", "", entry.description),
                    "item_name": item.name,
                    "amount": _claim_item_total(item.unit_price, item.tax_rate, item.quantity),
                    "quantity": item.quantity,
                    "unit": item.unit,
                    "unit_price": item.unit_price,
                    "budget_name": budget.name,
                    "budget_item": node.name,
                    "payment_method": entry.payment_method,
                    "status": public_status,
                    "evidence": [],
                },
            )
        )

    direct_rows = (
        await db.execute(
            select(
                FinanceBudgetExpense,
                FinanceBudgetExpenseItem,
                FinanceBudget,
                FinanceBudgetAllocation,
                FinanceBudgetNode,
            )
            .join(
                FinanceBudgetExpenseItem,
                FinanceBudgetExpenseItem.expense_id == FinanceBudgetExpense.id,
            )
            .join(FinanceBudget, FinanceBudget.id == FinanceBudgetExpense.budget_id)
            .join(
                FinanceBudgetAllocation,
                FinanceBudgetAllocation.id == FinanceBudgetExpense.allocation_id,
            )
            .join(FinanceBudgetNode, FinanceBudgetNode.id == FinanceBudgetAllocation.node_id)
            .join(
                FinanceBudgetSubmission,
                FinanceBudgetSubmission.id == FinanceBudgetAllocation.submission_id,
            )
            .where(
                FinanceBudget.is_public.is_(True),
                FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
                approved_initial,
            )
            .order_by(
                FinanceBudgetExpense.entry_date.desc(),
                FinanceBudgetExpense.created_at.desc(),
                FinanceBudgetExpenseItem.sort_order,
            )
            .limit(limit)
        )
    ).all()
    direct_expense_ids = list({expense.id for expense, _, _, _, _ in direct_rows})
    direct_evidence_rows = (
        list(
            (
                await db.execute(
                    select(FinanceBudgetExpenseEvidence)
                    .where(FinanceBudgetExpenseEvidence.expense_id.in_(direct_expense_ids))
                    .order_by(FinanceBudgetExpenseEvidence.created_at)
                )
            ).scalars()
        )
        if direct_expense_ids
        else []
    )
    evidence_by_expense: dict[uuid.UUID, list[FinanceBudgetExpenseEvidence]] = {}
    for evidence in direct_evidence_rows:
        evidence_by_expense.setdefault(evidence.expense_id, []).append(evidence)
    for expense, item, budget, _, node in direct_rows:
        public_expenses.append(
            (
                expense.entry_date,
                expense.created_at,
                {
                    "id": item.id,
                    "budget_id": budget.id,
                    "entry_date": expense.entry_date,
                    "purpose": expense.purpose,
                    "item_name": item.name,
                    "amount": _claim_item_total(item.unit_price, item.tax_rate, item.quantity),
                    "quantity": item.quantity,
                    "unit": item.unit,
                    "unit_price": item.unit_price,
                    "budget_name": budget.name,
                    "budget_item": node.name,
                    "payment_method": ExpensePaymentMethod.DIRECT,
                    "status": "spent",
                    "evidence": [
                        {
                            "id": evidence.id,
                            "filename": evidence.filename,
                            "url": (
                                f"/finance/public/budgets/{budget.id}/expenses/{expense.id}/"
                                f"evidence/{evidence.id}"
                            ),
                        }
                        for evidence in evidence_by_expense.get(expense.id, [])
                    ],
                },
            )
        )
    public_expenses.sort(key=lambda row: (row[0], row[1]), reverse=True)
    return [row[2] for row in public_expenses[:limit]]


async def public_expense_claim_detail(
    db: AsyncSession, budget_id: uuid.UUID, entry_id: uuid.UUID
) -> dict:
    budget = await db.get(FinanceBudget, budget_id)
    if not budget or not budget.is_public:
        raise HTTPException(404, "公開支出不存在")

    approved_initial = await db.scalar(
        select(FinanceBudgetSubmission.id).where(
            FinanceBudgetSubmission.budget_id == budget.id,
            FinanceBudgetSubmission.kind == BudgetSubmissionKind.INITIAL,
            FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
        )
    )
    if not approved_initial:
        raise HTTPException(404, "公開支出不存在")

    entry = await db.scalar(
        select(JournalEntry).where(
            JournalEntry.id == entry_id,
            JournalEntry.ledger_id == budget.ledger_id,
            JournalEntry.source_type == "expense_claim",
            JournalEntry.status == JournalStatus.POSTED,
            JournalEntry.claim_status == ExpenseClaimStatus.COMPLETED,
            JournalEntry.budget_included.is_(True),
        )
    )
    if not entry:
        raise HTTPException(404, "公開支出不存在")

    item_rows = (
        await db.execute(
            select(ExpenseClaimItem, FinanceBudgetNode)
            .join(FinanceBudgetNode, FinanceBudgetNode.id == ExpenseClaimItem.budget_node_id)
            .where(
                ExpenseClaimItem.journal_entry_id == entry.id,
                FinanceBudgetNode.budget_id == budget.id,
                FinanceBudgetNode.is_active.is_(True),
            )
            .order_by(ExpenseClaimItem.created_at, ExpenseClaimItem.id)
        )
    ).all()
    if not item_rows:
        raise HTTPException(404, "公開支出不存在")

    item_node_ids = {node.id for _, node in item_rows}
    approved_node_ids = set(
        (
            await db.execute(
                select(FinanceBudgetAllocation.node_id)
                .join(FinanceBudgetSubmission)
                .where(
                    FinanceBudgetAllocation.node_id.in_(item_node_ids),
                    FinanceBudgetSubmission.budget_id == budget.id,
                    FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
                )
            )
        ).scalars()
    )
    item_rows = [(item, node) for item, node in item_rows if node.id in approved_node_ids]
    if not item_rows:
        raise HTTPException(404, "公開支出不存在")

    item_ids = [item.id for item, _ in item_rows]
    evidence_rows = list(
        (
            await db.execute(
                select(ExpenseClaimItemEvidence)
                .where(ExpenseClaimItemEvidence.item_id.in_(item_ids))
                .order_by(ExpenseClaimItemEvidence.created_at, ExpenseClaimItemEvidence.id)
            )
        ).scalars()
    )
    evidence_by_item: dict[uuid.UUID, list[ExpenseClaimItemEvidence]] = {}
    for evidence in evidence_rows:
        evidence_by_item.setdefault(evidence.item_id, []).append(evidence)

    reporter_name = await db.scalar(select(User.display_name).where(User.id == entry.created_by_id))
    department_name = (
        await db.scalar(select(Org.name).where(Org.id == entry.proposing_org_id))
        if entry.proposing_org_id
        else None
    )
    items = [
        {
            "id": item.id,
            "name": item.name,
            "unit_price": item.unit_price,
            "tax_rate": item.tax_rate,
            "quantity": item.quantity,
            "unit": item.unit,
            "budget_item": node.name,
            "amount": _claim_item_total(item.unit_price, item.tax_rate, item.quantity),
            "evidence": [
                {
                    "id": evidence.id,
                    "filename": evidence.filename,
                    "url": (
                        f"/finance/public/budgets/{budget.id}/expense-claims/{entry.id}/"
                        f"evidence/{evidence.id}"
                    ),
                }
                for evidence in evidence_by_item.get(item.id, [])
            ],
        }
        for item, node in item_rows
    ]
    supplemental_evidence = (
        [
            {
                "id": entry.id,
                "filename": "報帳附件",
                "url": (
                    f"/finance/public/budgets/{budget.id}/expense-claims/{entry.id}/"
                    f"evidence/{entry.id}"
                ),
            }
        ]
        if entry.evidence_url
        else []
    )
    return {
        "id": entry.id,
        "budget_id": budget.id,
        "budget_name": budget.name,
        "entry_date": entry.entry_date,
        "purpose": re.sub(r"^報帳｜|（\d+ 項）$", "", entry.description),
        "department_name": department_name,
        "reporter_name": reporter_name or "已停用使用者",
        "reported_at": entry.created_at,
        "paid_at": entry.payment_at,
        "payment_method": entry.payment_method,
        "total_amount": sum(item["amount"] for item in items),
        "items": items,
        "supplemental_evidence": supplemental_evidence,
    }


async def public_budget_period_totals(db: AsyncSession, budget_id: uuid.UUID) -> dict:
    """Only expose ledger totals for the period of a published, approved budget."""
    budget = await db.get(FinanceBudget, budget_id)
    if not budget or not budget.is_public:
        raise HTTPException(404, "公開預算不存在")
    approved_initial = await db.scalar(
        select(FinanceBudgetSubmission.id).where(
            FinanceBudgetSubmission.budget_id == budget.id,
            FinanceBudgetSubmission.kind == BudgetSubmissionKind.INITIAL,
            FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
        )
    )
    if not approved_initial:
        raise HTTPException(404, "公開預算不存在")

    income, expenses = (
        await db.execute(
            select(
                func.coalesce(
                    func.sum(
                        case(
                            (
                                ChartAccount.account_type == FinanceAccountType.REVENUE,
                                JournalLine.credit - JournalLine.debit,
                            ),
                            else_=0,
                        )
                    ),
                    0,
                ),
                func.coalesce(
                    func.sum(
                        case(
                            (
                                ChartAccount.account_type == FinanceAccountType.EXPENSE,
                                JournalLine.debit - JournalLine.credit,
                            ),
                            else_=0,
                        )
                    ),
                    0,
                ),
            )
            .select_from(JournalLine)
            .join(ChartAccount, ChartAccount.id == JournalLine.account_id)
            .join(JournalEntry, JournalEntry.id == JournalLine.entry_id)
            .where(
                JournalEntry.ledger_id == budget.ledger_id,
                JournalEntry.period_id == budget.period_id,
                JournalEntry.status == JournalStatus.POSTED,
            )
        )
    ).one()
    direct_expenses = await db.scalar(
        select(func.coalesce(func.sum(FinanceBudgetExpense.total_amount), 0)).where(
            FinanceBudgetExpense.budget_id == budget.id
        )
    )
    return {
        "income_total": int(income),
        "expense_total": int(expenses) + int(direct_expenses or 0),
    }


async def public_budget_detail(
    db: AsyncSession, budget_id: uuid.UUID, review_submission_id: uuid.UUID | None = None
) -> tuple[dict, FiscalPeriod, FinanceBudgetSubmission | None]:
    budget = await db.get(FinanceBudget, budget_id)
    if not budget:
        raise HTTPException(404, "公開預算不存在")
    period = await db.get(FiscalPeriod, budget.period_id)
    if not period:
        raise HTTPException(404, "會計期間不存在")
    if review_submission_id:
        submission = await db.scalar(
            select(FinanceBudgetSubmission).where(
                FinanceBudgetSubmission.id == review_submission_id,
                FinanceBudgetSubmission.budget_id == budget.id,
                FinanceBudgetSubmission.is_council_review_public,
                FinanceBudgetSubmission.status.in_(
                    [
                        BudgetSubmissionStatus.DRAFT,
                        BudgetSubmissionStatus.SUBMITTED,
                        BudgetSubmissionStatus.RETURNED,
                    ]
                ),
            )
        )
        if not submission:
            raise HTTPException(404, "公開預算不存在")
        return await budget_detail(db, budget), period, submission

    if not budget.is_public:
        raise HTTPException(404, "公開預算不存在")
    approved_initial = await db.scalar(
        select(FinanceBudgetSubmission.id).where(
            FinanceBudgetSubmission.budget_id == budget.id,
            FinanceBudgetSubmission.kind == "initial",
            FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
        )
    )
    if not approved_initial:
        raise HTTPException(404, "公開預算不存在")
    detail = await budget_detail(db, budget)
    for expense in detail["expenses"]:
        for evidence in expense["evidence"]:
            evidence["url"] = (
                f"/finance/public/budgets/{budget.id}/expenses/{expense['id']}/"
                f"evidence/{evidence['id']}"
            )
    return detail, period, None


async def settlement_report(db: AsyncSession, ledger_id: uuid.UUID, period_id: uuid.UUID) -> dict:
    period = await db.get(FiscalPeriod, period_id)
    if not period or period.ledger_id != ledger_id:
        raise HTTPException(404, "會計期間不存在")
    budget = await db.scalar(
        select(FinanceBudget).where(
            FinanceBudget.ledger_id == ledger_id, FinanceBudget.period_id == period_id
        )
    )
    if not budget:
        return {
            "period_id": period.id,
            "period_name": period.name,
            "budgeted_total": 0,
            "settled_total": 0,
            "unsettled_claim_count": 0,
            "lines": [],
        }
    detail = await budget_detail(db, budget)
    settled_rows = (
        await db.execute(
            select(
                ExpenseClaimItem.budget_node_id,
                func.coalesce(
                    func.sum(
                        func.round(
                            ExpenseClaimItem.unit_price
                            * ((100 + ExpenseClaimItem.tax_rate) / 100)
                            * ExpenseClaimItem.quantity
                        )
                    ),
                    0,
                ),
            )
            .join(JournalEntry, JournalEntry.id == ExpenseClaimItem.journal_entry_id)
            .where(
                ExpenseClaimItem.budget_node_id.is_not(None),
                JournalEntry.ledger_id == ledger_id,
                JournalEntry.period_id == period_id,
                JournalEntry.status == JournalStatus.POSTED,
                JournalEntry.claim_status == ExpenseClaimStatus.COMPLETED,
            )
            .group_by(ExpenseClaimItem.budget_node_id)
        )
    ).all()
    settled = {row[0]: int(row[1]) for row in settled_rows if row[0] is not None}
    direct_settled_rows = (
        await db.execute(
            select(FinanceBudgetAllocation.node_id, func.sum(FinanceBudgetExpense.total_amount))
            .join(
                FinanceBudgetExpense,
                FinanceBudgetExpense.allocation_id == FinanceBudgetAllocation.id,
            )
            .join(
                FinanceBudgetSubmission,
                FinanceBudgetSubmission.id == FinanceBudgetAllocation.submission_id,
            )
            .where(
                FinanceBudgetExpense.budget_id == budget.id,
                FinanceBudgetSubmission.status == BudgetSubmissionStatus.APPROVED,
            )
            .group_by(FinanceBudgetAllocation.node_id)
        )
    ).all()
    for node_id, amount in direct_settled_rows:
        settled[node_id] = settled.get(node_id, 0) + int(amount or 0)
    unsettled_claim_count = await db.scalar(
        select(func.count())
        .select_from(JournalEntry)
        .where(
            JournalEntry.ledger_id == ledger_id,
            JournalEntry.period_id == period_id,
            JournalEntry.source_type == "expense_claim",
            JournalEntry.status == JournalStatus.POSTED,
            JournalEntry.claim_status != ExpenseClaimStatus.COMPLETED,
        )
    )
    lines = [
        {
            "node_id": node["id"],
            "name": node["name"],
            "budgeted_amount": node["allocated_amount"],
            "settled_amount": settled.get(node["id"], 0),
            "difference_amount": node["allocated_amount"] - settled.get(node["id"], 0),
        }
        for node in detail["nodes"]
        if not any(child["parent_id"] == node["id"] for child in detail["nodes"])
    ]
    return {
        "period_id": period.id,
        "period_name": period.name,
        "budgeted_total": sum(line["budgeted_amount"] for line in lines),
        "settled_total": sum(line["settled_amount"] for line in lines),
        "unsettled_claim_count": int(unsettled_claim_count or 0),
        "lines": lines,
    }
