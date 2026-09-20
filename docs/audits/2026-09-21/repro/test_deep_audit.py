import json
import uuid
from datetime import UTC, date, datetime, timedelta
from types import SimpleNamespace

import pytest
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker

from api.core.config import settings
from api.core.security import create_access_token
from api.models.org import Org, Position, UserPosition
from api.models.survey import (
    QuestionType,
    Survey,
    SurveyAnswer,
    SurveyQuestion,
    SurveyResponse,
    SurveyStatus,
)
from api.schemas.survey import AnswerSubmit, SurveySubmit
from api.services import survey as svc


async def make_survey(db, user, **kwargs):
    org = Org(name="audit-test-org")
    db.add(org)
    await db.flush()
    survey = Survey(
        title="audit-test",
        org_id=org.id,
        created_by=user.id,
        status=SurveyStatus.OPEN,
        questions=[],
        **kwargs,
    )
    db.add(survey)
    await db.flush()
    return org, survey


async def test_anonymous_response_audit_reidentifies_user(db_session, member_user, monkeypatch):
    from api.models.audit_log import AuditLog
    from api.routers import survey as router

    monkeypatch.setattr(router, "get_posthog_client", lambda: None)
    _, survey = await make_survey(db_session, member_user, is_anonymous=True, is_public=True)
    q = SurveyQuestion(
        survey_id=survey.id,
        question_text="opinion",
        question_type=QuestionType.TEXT,
        is_required=True,
        order_index=0,
    )
    db_session.add(q)
    await db_session.flush()
    response = await router.submit_response(
        str(survey.id),
        SurveySubmit(
            answers=[AnswerSubmit(question_id=q.id, answer_text="private answer")],
            anon_token=uuid.uuid4().hex,
        ),
        db_session,
        member_user,
    )
    log = await db_session.scalar(
        select(AuditLog).where(
            AuditLog.action == "survey.response_submit", AuditLog.entity_id == str(response.id)
        )
    )
    print(
        "EVIDENCE anonymity",
        json.dumps(
            {
                "respondent_id": response.respondent_id,
                "audit_actor_matches_user": log.actor_id == str(member_user.id),
                "audit_email_matches": log.actor_email == member_user.email,
                "audit_entity_matches_response": log.entity_id == str(response.id),
            }
        ),
    )
    assert response.respondent_id is None and log.actor_id == str(member_user.id)


async def test_required_empty_invalid_choice_and_anonymous_repeats_accepted(
    db_session, member_user
):
    _, survey = await make_survey(
        db_session, member_user, is_anonymous=True, is_public=True, allow_multiple=False
    )
    q1 = SurveyQuestion(
        survey_id=survey.id,
        question_text="required",
        question_type=QuestionType.TEXT,
        is_required=True,
        order_index=0,
    )
    q2 = SurveyQuestion(
        survey_id=survey.id,
        question_text="choice",
        question_type=QuestionType.SINGLE,
        is_required=True,
        options_json='["A","B"]',
        order_index=1,
    )
    db_session.add_all([q1, q2])
    await db_session.flush()
    for _ in range(2):
        await svc.submit_response(
            db_session,
            survey,
            respondent_id=None,
            data=SurveySubmit(
                answers=[
                    AnswerSubmit(question_id=q1.id),
                    AnswerSubmit(question_id=q2.id, answer_options=["NOT-AN-OPTION"]),
                ]
            ),
        )
    count = await db_session.scalar(
        select(func.count(SurveyResponse.id)).where(SurveyResponse.survey_id == survey.id)
    )
    answers = (await db_session.scalars(select(SurveyAnswer.answer_text))).all()
    print(
        "EVIDENCE survey validation",
        json.dumps(
            {"allow_multiple": False, "responses_without_token": count, "saved_answers": answers}
        ),
    )
    assert count == 2 and None in answers and "NOT-AN-OPTION" in answers


async def test_nonanonymous_allow_multiple_conflicts_with_unique_constraint(
    db_session, member_user
):
    from sqlalchemy.exc import IntegrityError

    _, survey = await make_survey(
        db_session, member_user, is_anonymous=False, is_public=False, allow_multiple=True
    )
    q = SurveyQuestion(
        survey_id=survey.id,
        question_text="opinion",
        question_type=QuestionType.TEXT,
        is_required=True,
        order_index=0,
    )
    db_session.add(q)
    await db_session.flush()
    data = SurveySubmit(answers=[AnswerSubmit(question_id=q.id, answer_text="yes")])
    await svc.submit_response(db_session, survey, respondent_id=member_user.id, data=data)
    with pytest.raises(IntegrityError):
        async with db_session.begin_nested():
            await svc.submit_response(db_session, survey, respondent_id=member_user.id, data=data)
    print("EVIDENCE allow_multiple=true second response rejected by DB unique constraint")


async def test_expired_tenure_still_allows_restricted_survey_and_ws(
    db_session, member_user, monkeypatch
):
    from api.routers import ws

    org, survey = await make_survey(db_session, member_user, is_anonymous=False, is_public=False)
    survey.allowed_org_ids_json = json.dumps([str(org.id)])
    pos = Position(name="expired", org_id=org.id)
    db_session.add(pos)
    await db_session.flush()
    db_session.add(
        UserPosition(
            user_id=member_user.id,
            position_id=pos.id,
            start_date=date(2020, 1, 1),
            end_date=date(2020, 12, 31),
        )
    )
    await db_session.flush()
    await svc.check_survey_access(db_session, survey, member_user)

    async def no_cached_permissions(*a, **k):
        return []

    monkeypatch.setattr("api.core.cache.cache_get", no_cached_permissions)
    monkeypatch.setattr(
        ws,
        "AsyncSessionLocal",
        async_sessionmaker(
            db_session.bind, expire_on_commit=False, join_transaction_mode="create_savepoint"
        ),
    )
    await ws._assert_room_access("org:" + str(org.id), str(member_user.id))
    print("EVIDENCE expired tenure 2020-12-31 accepted: restricted survey + organization WS")


async def test_revoked_session_ws_accepts_http_rejects(db_session, member_user, monkeypatch):
    from api.dependencies import auth
    from api.models.user_session import UserSession
    from api.routers import ws

    now = datetime.now(UTC)
    session = UserSession(
        user_id=member_user.id,
        refresh_jti_hash=uuid.uuid4().hex,
        auth_time=now,
        last_seen_at=now,
        rotated_at=now,
        expires_at=now + timedelta(days=1),
        absolute_expires_at=now + timedelta(days=7),
        revoked_at=now,
        revoked_reason="logout",
    )
    db_session.add(session)
    await db_session.flush()
    token = create_access_token(str(member_user.id), session_id=str(session.id))

    async def absent(*a, **k):
        return False

    for mod in [ws, auth]:
        monkeypatch.setattr(mod, "is_blacklisted", absent)
        monkeypatch.setattr(mod, "is_session_revoked", absent)

    async def close(*a, **k):
        raise AssertionError("unexpected close")

    socket = SimpleNamespace(
        headers={}, cookies={settings.ACCESS_TOKEN_COOKIE_NAME: token}, close=close
    )
    ws_result = await ws._authenticate_ws(socket)
    http_result = await auth._user_from_access_token(token, db_session)
    print(
        "EVIDENCE revoked session",
        json.dumps(
            {"ws_accepted": ws_result is not None, "http_accepted": http_result is not None}
        ),
    )
    assert ws_result is not None and http_result is None


async def test_meal_platform_list_schema_rejects_real_row(db_session, member_user):
    from pydantic import ValidationError

    from api.models.meal import MealOrder, MealVendor
    from api.schemas.meal import MealOrderListItem

    org = Org(name="audit-meal")
    db_session.add(org)
    await db_session.flush()
    vendor = MealVendor(name="audit-vendor", org_id=org.id, created_by=member_user.id)
    db_session.add(vendor)
    await db_session.flush()
    order = MealOrder(
        user_id=member_user.id,
        vendor_id=vendor.id,
        schedule_id=None,
        serial_number="AUDIT-PLATFORM",
        pickup_code="12345",
        total_price=50,
    )
    db_session.add(order)
    await db_session.flush()
    with pytest.raises(ValidationError) as e:
        MealOrderListItem.model_validate(order)
    print("EVIDENCE platform list schema", str(e.value))


async def test_legacy_meal_duplicate_line_rejected_by_db(db_session, member_user):
    from api.models.meal import MealVendor, MenuItem, MenuSchedule
    from api.schemas.meal import MealOrderCreate, MealOrderItemCreate
    from api.services.meal import create_meal_order

    org = Org(name="audit-meal-stock")
    db_session.add(org)
    await db_session.flush()
    vendor = MealVendor(name="audit-stock", org_id=org.id, created_by=member_user.id)
    db_session.add(vendor)
    await db_session.flush()
    schedule = MenuSchedule(
        vendor_id=vendor.id,
        date=date.today(),
        order_deadline=datetime.now(UTC) + timedelta(days=1),
        created_by=member_user.id,
    )
    db_session.add(schedule)
    await db_session.flush()
    item = MenuItem(
        schedule_id=schedule.id, name="limited", price=10, max_quantity=5, is_available=True
    )
    db_session.add(item)
    await db_session.flush()
    from sqlalchemy.exc import IntegrityError

    with pytest.raises(IntegrityError):
        async with db_session.begin_nested():
            await create_meal_order(
                db_session,
                user_id=member_user.id,
                data=MealOrderCreate(
                    schedule_id=schedule.id,
                    items=[
                        MealOrderItemCreate(menu_item_id=item.id, quantity=4),
                        MealOrderItemCreate(menu_item_id=item.id, quantity=4),
                    ],
                ),
            )
    print(
        "ELIMINATED duplicate menu item oversell: database unique constraint rejects duplicate lines"
    )


async def test_meal_platform_list_http_fails(db_session, member_user, authed_client_factory):
    from fastapi.exceptions import ResponseValidationError

    from api.models.meal import MealOrder, MealVendor

    org = Org(name="audit-http")
    db_session.add(org)
    await db_session.flush()
    vendor = MealVendor(name="audit-http", org_id=org.id, created_by=member_user.id)
    db_session.add(vendor)
    await db_session.flush()
    db_session.add(
        MealOrder(
            user_id=member_user.id,
            vendor_id=vendor.id,
            schedule_id=None,
            serial_number="AUDIT-HTTP",
            pickup_code="12346",
            total_price=50,
        )
    )
    await db_session.flush()
    with pytest.raises(ResponseValidationError) as exc:
        await authed_client_factory(member_user).get("/meal/orders")
    print("EVIDENCE real HTTP GET /meal/orders", str(exc.value))


async def test_meal_manager_can_read_other_org_order_despite_list_scope(
    db_session, member_user, make_user, authed_client_factory
):
    from api.models.meal import MealOrder, MealVendor
    from api.models.org import Permission

    a = Org(name="managed-A")
    b = Org(name="unmanaged-B")
    db_session.add_all([a, b])
    await db_session.flush()
    p = Position(name="meal-manager", org_id=a.id)
    db_session.add(p)
    await db_session.flush()
    db_session.add_all(
        [
            UserPosition(user_id=member_user.id, position_id=p.id, start_date=date(2020, 1, 1)),
            Permission(position_id=p.id, code="meal:manage"),
        ]
    )
    target = await make_user()
    vendor = MealVendor(name="vendor-B", org_id=b.id, created_by=target.id)
    db_session.add(vendor)
    await db_session.flush()
    order = MealOrder(
        user_id=target.id,
        vendor_id=vendor.id,
        schedule_id=None,
        serial_number="AUDIT-CROSS-ORG",
        pickup_code="15376",
        total_price=150,
        notes="private notes",
    )
    db_session.add(order)
    await db_session.flush()
    ac = authed_client_factory(member_user)
    list_res = await ac.get(
        "/meal/orders", params={"my_only": "false", "vendor_id": str(vendor.id)}
    )
    detail = await ac.get("/meal/orders/" + str(order.id))
    print(
        "EVIDENCE cross-org meal",
        json.dumps(
            {
                "list_status": list_res.status_code,
                "detail_status": detail.status_code,
                "pickup_code_exposed": detail.json().get("pickup_code") == order.pickup_code,
                "other_user_id_exposed": detail.json().get("user_id") == str(target.id),
            }
        ),
    )
    assert list_res.status_code == 403 and detail.status_code == 200


async def test_expired_platform_meal_can_be_cancelled(db_session, member_user):
    from api.models.meal import (
        MealOrder,
        MealOrderStatus,
        MealPickupSlot,
        MealProduct,
        MealProductAvailability,
        MealVendor,
    )
    from api.services.meal import cancel_meal_order

    org = Org(name="deadline-org")
    db_session.add(org)
    await db_session.flush()
    v = MealVendor(name="deadline-vendor", org_id=org.id, created_by=member_user.id)
    db_session.add(v)
    await db_session.flush()
    p = MealProduct(vendor_id=v.id, name="lunch", price=50)
    db_session.add(p)
    await db_session.flush()
    a = MealProductAvailability(
        product_id=p.id, vendor_id=v.id, service_date=date.today(), price=50
    )
    db_session.add(a)
    await db_session.flush()
    now = datetime.now(UTC)
    slot = MealPickupSlot(
        availability_id=a.id,
        label="noon",
        pickup_start=now - timedelta(hours=1),
        pickup_end=now + timedelta(hours=1),
        order_deadline=now - timedelta(hours=3),
    )
    db_session.add(slot)
    await db_session.flush()
    order = MealOrder(
        user_id=member_user.id,
        vendor_id=v.id,
        schedule_id=None,
        availability_id=a.id,
        pickup_slot_id=slot.id,
        serial_number="AUDIT-DEADLINE",
        pickup_code="14346",
        total_price=50,
        status=MealOrderStatus.CONFIRMED,
    )
    db_session.add(order)
    await db_session.flush()
    await cancel_meal_order(db_session, order, requested_by=member_user.id)
    print(
        "EVIDENCE past deadline platform meal",
        json.dumps(
            {"deadline_hours_ago": 3, "original_status": "confirmed", "result_status": order.status}
        ),
    )
    assert order.status == MealOrderStatus.CANCELLED


async def test_survey_excel_formula_and_event_loop_blocking(db_session, member_user):
    import asyncio
    import io
    import time

    from openpyxl import load_workbook
    from sqlalchemy import insert

    _, survey = await make_survey(
        db_session, member_user, is_anonymous=True, is_public=True, allow_multiple=True
    )
    questions = [
        SurveyQuestion(
            survey_id=survey.id,
            question_text=f"Question {i}",
            question_type=QuestionType.TEXT,
            is_required=True,
            order_index=i,
        )
        for i in range(8)
    ]
    db_session.add_all(questions)
    await db_session.flush()
    ids = [uuid.uuid4() for _ in range(2500)]
    await db_session.execute(
        insert(SurveyResponse),
        [{"id": r, "survey_id": survey.id, "submitted_at": datetime.now(UTC)} for r in ids],
    )
    await db_session.execute(
        insert(SurveyAnswer),
        [
            {
                "id": uuid.uuid4(),
                "response_id": r,
                "question_id": q.id,
                "answer_text": "=1+1" if j == 0 and i == 0 else "audit feedback " + str(j),
            }
            for j, r in enumerate(ids)
            for i, q in enumerate(questions)
        ],
    )
    gaps = []
    running = True

    async def heartbeat():
        last = time.perf_counter()
        while running:
            await asyncio.sleep(0.01)
            now = time.perf_counter()
            gaps.append((now - last) * 1000)
            last = now

    task = asyncio.create_task(heartbeat())
    await asyncio.sleep(0.03)
    start = time.perf_counter()
    data = await svc.build_survey_export(db_session, survey)
    elapsed = time.perf_counter() - start
    await asyncio.sleep(0.03)
    running = False
    await task
    wb = load_workbook(io.BytesIO(data), data_only=False)
    formula_cells = [
        {"cell": c.coordinate, "type": c.data_type, "value": c.value}
        for row in wb["回應明細"]
        for c in row
        if c.data_type == "f"
    ]
    print(
        "EVIDENCE export",
        json.dumps(
            {
                "rows": 2500,
                "questions": 8,
                "seconds": round(elapsed, 3),
                "max_10ms_heartbeat_gap_ms": round(max(gaps), 1),
                "xlsx_bytes": len(data),
                "formula_cells": formula_cells,
            }
        ),
    )
    assert formula_cells and formula_cells[0]["value"] == "=1+1"


async def test_calendar_projection_discloses_restricted_document(
    db_session, member_user, make_user
):
    from api.models.document import Document, DocumentStatus, DocumentVisibility
    from api.schemas.calendar import CalendarEventListItem
    from api.services.calendar import list_events
    from api.services.coordination import _project_documents
    from api.services.document import check_document_access

    creator = await make_user()
    org = Org(name="audit-calendar-org")
    db_session.add(org)
    await db_session.flush()
    pos = Position(name="ordinary org member", org_id=org.id)
    db_session.add(pos)
    await db_session.flush()
    db_session.add(
        UserPosition(user_id=member_user.id, position_id=pos.id, start_date=date(2020, 1, 1))
    )
    doc = Document(
        serial_number="AUDIT-" + uuid.uuid4().hex[:10],
        org_id=org.id,
        created_by=creator.id,
        title="Synthetic restricted case",
        subject="Synthetic private case subject",
        visibility_level=DocumentVisibility.SUBJECT_ONLY,
        status=DocumentStatus.DRAFT,
        due_date=datetime.now(UTC) + timedelta(days=1),
    )
    db_session.add(doc)
    await db_session.flush()
    await db_session.refresh(doc, attribute_names=["approvals"])
    direct_access = await check_document_access(db_session, doc, member_user.id)
    await _project_documents(db_session, None, None)
    visible = await list_events(db_session, user=member_user, permission_codes=frozenset())
    leaked = [
        CalendarEventListItem.model_validate(e).model_dump(mode="json")
        for e in visible
        if e.source_id == doc.id
    ]
    print(
        "EVIDENCE calendar source permissions",
        json.dumps(
            {
                "direct_document_allowed": direct_access,
                "calendar_exposes_title": bool(leaked),
                "calendar_description": leaked[0]["description"] if leaked else None,
                "source_visibility": doc.visibility_level,
            }
        ),
    )
    assert direct_access is False and leaked and leaked[0]["description"] == doc.subject


async def test_calendar_projection_race_breaks_transaction():
    import asyncio

    from conftest import _schema_engine
    from sqlalchemy.ext.asyncio import AsyncSession

    from api.models.user import User
    from api.services.coordination import _upsert_projection

    factory = async_sessionmaker(_schema_engine, expire_on_commit=False)
    async with factory() as seed:
        actor = User(email=f"audit-race-{uuid.uuid4().hex}@example.test", display_name="Audit race")
        seed.add(actor)
        await seed.commit()
        actor_id = actor.id
    count = 0
    both_read = asyncio.Event()
    source_id = uuid.uuid4()

    class ScheduledSession(AsyncSession):
        async def scalar(self, *args, **kwargs):
            nonlocal count
            result = await super().scalar(*args, **kwargs)
            if not getattr(self, "audit_first_read", False):
                self.audit_first_read = True
                count += 1
                if count == 2:
                    both_read.set()
                await asyncio.wait_for(both_read.wait(), 5)
            return result

    concurrent_factory = async_sessionmaker(
        _schema_engine, class_=ScheduledSession, expire_on_commit=False
    )

    async def run_one():
        async with concurrent_factory() as session:
            try:
                await _upsert_projection(
                    session,
                    source_module="audit-race",
                    source_id=source_id,
                    source_key="due",
                    org_id=None,
                    title="synthetic race",
                    starts_at=datetime.now(UTC),
                    created_by=actor_id,
                    href="/audit",
                )
                await session.commit()
                return "committed"
            except Exception as exc:
                return type(exc).__name__

    outcomes = await asyncio.wait_for(asyncio.gather(run_one(), run_one()), 15)
    print("EVIDENCE calendar simultaneous first projection", json.dumps(outcomes))
    assert sorted(outcomes) == ["PendingRollbackError", "committed"]


async def test_private_document_anonymous_404_owner_200(db_session, member_user):
    from fastapi import HTTPException

    from api.models.document import Document, DocumentStatus, DocumentVisibility
    from api.routers.documents import get_document

    org = Org(name="audit-private-page")
    db_session.add(org)
    await db_session.flush()
    doc = Document(
        serial_number="AUDIT-SSR",
        org_id=org.id,
        created_by=member_user.id,
        title="Private SSR case",
        status=DocumentStatus.DRAFT,
        visibility_level=DocumentVisibility.SUBJECT_ONLY,
    )
    db_session.add(doc)
    await db_session.flush()
    with pytest.raises(HTTPException) as exc:
        await get_document(str(doc.id), db_session, None)
    owner_result = await get_document(str(doc.id), db_session, member_user)
    print(
        "EVIDENCE private document API",
        json.dumps(
            {
                "anonymous_status": exc.value.status_code,
                "owner_gets_document": owner_result.id == doc.id,
            }
        ),
    )
    assert exc.value.status_code == 404 and owner_result.id == doc.id
